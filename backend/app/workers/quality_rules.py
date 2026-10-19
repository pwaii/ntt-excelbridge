"""Deterministic structural quality checks (§9)."""

from __future__ import annotations

from collections import defaultdict
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.workers.inspect_workbook import inspect_bytes, normalize_text


def _issue(issue_id: str, kind: str, message: str, **details: Any) -> dict[str, Any]:
    return {"issue_id": issue_id, "kind": kind, "message": message, "recommended_action": details.pop("recommended_action", "manual_review"), **details}


def run_quality_rules(source: bytes, manifest: dict[str, Any] | None = None) -> dict[str, Any]:
    """Return stable issue records for duplicate/blank/type problems."""

    manifest = manifest or inspect_bytes(source)
    workbook = load_workbook(BytesIO(source), read_only=True, data_only=False)
    issues: list[dict[str, Any]] = []
    try:
        for sheet in manifest["sheets"]:
            worksheet = workbook[sheet["name"]]
            rows = list(worksheet.iter_rows(values_only=True))
            header_row = sheet["header_row"]
            headers = sheet["headers"]
            data_rows = [tuple(row[i] if i < len(row) else None for i in range(len(headers))) for row in rows[header_row:]]
            duplicates: dict[str, list[int]] = defaultdict(list)
            for row_number, row in enumerate(data_rows, start=header_row + 1):
                if all(normalize_text(cell) == "" for cell in row):
                    issues.append(_issue(f"BLANK_ROW#{sheet['name']}#{row_number}", "blank_row", "row is empty after normalization", sheet=sheet["name"], row_numbers=[row_number], recommended_action="delete"))
                    continue
                fingerprint = next(item["fingerprint"] for item in sheet["row_fingerprints"] if item["row_number"] == row_number)
                duplicates[fingerprint].append(row_number)
            for fingerprint, row_numbers in duplicates.items():
                if len(row_numbers) > 1:
                    issues.append(_issue(f"DUPLICATE#{sheet['name']}#{fingerprint[:16]}", "duplicate_rows", "rows have the same normalized values", sheet=sheet["name"], row_numbers=row_numbers, fingerprint=fingerprint, keep_row=row_numbers[0], recommended_action="keep_first_delete_later"))
            for column in sheet["columns"]:
                if column["blank_count"] == len(data_rows) and data_rows:
                    issues.append(_issue(f"BLANK_COLUMN#{sheet['name']}#{column['index']}", "blank_column", "column is empty for every data row", sheet=sheet["name"], column=column["name"], recommended_action="exclude_from_input"))
                distribution = column["type_distribution"]
                total = sum(distribution.values())
                if total < 2:
                    continue
                dominant_type, dominant_count = max(distribution.items(), key=lambda item: item[1])
                minority = total - dominant_count
                if dominant_count / total < 0.8 or minority / total > 0.2:
                    target_type = {
                        "decimal": "number",
                        "percentage": "number",
                        "identifier": "string",
                        "text": "string",
                        "formula": "string",
                    }.get(dominant_type, dominant_type)
                    issues.append(_issue(f"TYPE_MISMATCH#{sheet['name']}#{column['index']}", "type_mismatch", "column contains incompatible value types", sheet=sheet["name"], column=column["name"], type_distribution=distribution, proposed_target_type=target_type, recommended_action="manual_review"))
        return {"report_version": "1", "issue_count": len(issues), "issues": issues}
    finally:
        workbook.close()
