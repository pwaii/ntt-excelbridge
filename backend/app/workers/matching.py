"""Deterministic current-to-main matching used before Tool #5.

The matching layer deliberately runs before the model. It reduces the main
workbook to a small candidate set and gives every candidate a stable row ID.
Tool #5 can choose from those IDs, but it cannot invent a new one.
"""

from __future__ import annotations

from collections import defaultdict
from difflib import SequenceMatcher
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.workers.inspect_workbook import normalize_text


def normalize_key(value: Any) -> str:
    """Normalize a business key without changing the stored workbook value."""

    return "".join(normalize_text(value).casefold().split())


def load_canonical_rows(source: bytes, sheet_name: str | None = None) -> list[dict[str, Any]]:
    workbook = load_workbook(BytesIO(source), read_only=True, data_only=False)
    try:
        worksheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
        rows = list(worksheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [normalize_text(value) or f"Column_{i + 1}" for i, value in enumerate(rows[0])]
        return [{"row_id": f"{worksheet.title}!{number}", **{headers[i]: row[i] if i < len(row) else None for i in range(len(headers))}} for number, row in enumerate(rows[1:], start=2)]
    finally:
        workbook.close()


def exact_match_rows(current_rows: list[dict[str, Any]], main_rows: list[dict[str, Any]], key_column: str) -> list[dict[str, Any]]:
    buckets: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in main_rows:
        key = normalize_key(row.get(key_column))
        if key:
            buckets[key].append(row)
    results = []
    for current in current_rows:
        key = normalize_key(current.get(key_column))
        candidates = buckets.get(key, []) if key else []
        if len(candidates) == 1:
            match_type = "exact"
            selected = candidates[0]["row_id"]
        elif len(candidates) > 1:
            match_type = "ambiguous"
            selected = None
        else:
            match_type = "unmatched"
            selected = None
        results.append({"current_row_id": current.get("row_id"), "candidate_main_row_ids": [row["row_id"] for row in candidates], "selected_main_row_id": selected, "match_type": match_type})
    return results


def _key_similarity(left: str, right: str) -> float:
    if not left or not right:
        return 0.0
    return SequenceMatcher(None, left, right, autojunk=False).ratio()


def generate_candidates(
    current_rows: list[dict[str, Any]],
    main_rows: list[dict[str, Any]],
    key_column: str,
    max_candidates: int = 5,
) -> list[dict[str, Any]]:
    """Create bounded, deterministic candidates for each current row.

    Exact and normalized matches are surfaced directly. For unresolved rows,
    candidates are blocked by a shared first character or a similar key length,
    then ranked by key similarity. The output contains row IDs and scores but
    never asks the model to search the complete main sheet.
    """

    limit = max(1, min(int(max_candidates), 25))
    keyed_main = [
        (row, normalize_key(row.get(key_column)))
        for row in main_rows
        if normalize_key(row.get(key_column))
    ]
    exact_buckets: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row, key in keyed_main:
        exact_buckets[key].append(row)

    results: list[dict[str, Any]] = []
    for current in current_rows:
        current_key = normalize_key(current.get(key_column))
        exact = exact_buckets.get(current_key, []) if current_key else []
        if len(exact) == 1:
            results.append(
                {
                    "current_row_id": current.get("row_id"),
                    "candidate_main_row_ids": [exact[0]["row_id"]],
                    "candidate_scores": {exact[0]["row_id"]: 1.0},
                    "selected_main_row_id": exact[0]["row_id"],
                    "match_type": "exact",
                }
            )
            continue
        if len(exact) > 1:
            results.append(
                {
                    "current_row_id": current.get("row_id"),
                    "candidate_main_row_ids": [row["row_id"] for row in exact[:limit]],
                    "candidate_scores": {row["row_id"]: 1.0 for row in exact[:limit]},
                    "selected_main_row_id": None,
                    "match_type": "duplicate",
                }
            )
            continue

        first = current_key[:1]
        blocked = [
            (row, key, _key_similarity(current_key, key))
            for row, key in keyed_main
            if (first and key[:1] == first) or abs(len(key) - len(current_key)) <= 2
        ]
        ranked = sorted(blocked, key=lambda item: (-item[2], item[1], item[0]["row_id"]))[:limit]
        scores = {row["row_id"]: round(score, 6) for row, _, score in ranked}
        results.append(
            {
                "current_row_id": current.get("row_id"),
                "candidate_main_row_ids": [row["row_id"] for row, _, _ in ranked],
                "candidate_scores": scores,
                "selected_main_row_id": None,
                "match_type": "unmatched" if not ranked else "ambiguous",
            }
        )
    return results


def build_tool5_inputs(
    current_rows: list[dict[str, Any]],
    main_rows: list[dict[str, Any]],
    key_column: str,
    max_candidates: int = 5,
) -> dict[str, Any]:
    """Build bounded Tool #5 input from rows and candidate IDs."""

    candidate_sets = generate_candidates(current_rows, main_rows, key_column, max_candidates)
    main_by_id = {row["row_id"]: row for row in main_rows}
    current_by_id = {row["row_id"]: row for row in current_rows}
    unresolved = []
    bounded_candidate_rows: dict[str, dict[str, Any]] = {}
    for item in candidate_sets:
        if item["match_type"] == "exact":
            continue
        unresolved.append(
            {
                "current_row_id": item["current_row_id"],
                "values": current_by_id.get(item["current_row_id"], {}),
                "candidate_main_row_ids": item["candidate_main_row_ids"],
                "candidate_scores": item["candidate_scores"],
                "candidate_rows": [main_by_id[row_id] for row_id in item["candidate_main_row_ids"]],
            }
        )
        for row_id in item["candidate_main_row_ids"]:
            if row_id in main_by_id:
                bounded_candidate_rows[row_id] = main_by_id[row_id]
    return {
        "current_rows": unresolved,
        "candidate_rows": list(bounded_candidate_rows.values()),
        "match_key": key_column,
        "deterministic_matches": [item for item in candidate_sets if item["match_type"] == "exact"],
    }
