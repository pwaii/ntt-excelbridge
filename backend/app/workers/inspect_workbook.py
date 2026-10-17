"""Deterministic OpenPyXL inspection used before any AI call.

The worker deliberately emits metadata and bounded, shape-redacted samples.  It
never puts workbook bytes or an unbounded row array into workflow state.
"""

from __future__ import annotations

import hashlib
import json
import re
import unicodedata
from collections import Counter
from collections.abc import Iterable
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.ai.policy import redact_sample
from app.settings import settings

_DATE_RE = re.compile(r"^(?:\d{4}[-/]\d{1,2}[-/]\d{1,2}|\d{1,2}[/-]\d{1,2}[/-]\d{2,4})$")
_PERCENT_RE = re.compile(r"^-?\d+(?:\.\d+)?%$")
_IDENTIFIER_RE = re.compile(r"^[A-Za-z]*\d+[A-Za-z\d_-]*$")


def normalize_text(value: Any) -> str:
    """Normalize whitespace and empty values for fingerprints and comparisons."""

    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value))
    return " ".join(text.replace("\u00a0", " ").split()).strip()


def _type_name(value: Any) -> str:
    if value is None or normalize_text(value) == "":
        return "unknown"
    if isinstance(value, str) and value.startswith("="):
        return "formula"
    if isinstance(value, bool):
        return "bool"
    if isinstance(value, datetime | date):
        return "date"
    if isinstance(value, int) and not isinstance(value, bool):
        return "int"
    if isinstance(value, float):
        return "decimal"
    text = normalize_text(value)
    if _PERCENT_RE.match(text):
        return "percentage"
    if _DATE_RE.match(text):
        return "date"
    if _IDENTIFIER_RE.match(text) and any(c.isdigit() for c in text):
        return "identifier-like"
    if text.lower() in {"true", "false", "yes", "no", "y", "n", "1", "0"}:
        return "bool"
    return "text"


def _header_row(rows: list[tuple[Any, ...]]) -> int:
    for index, row in enumerate(rows[:20], start=1):
        non_empty = [cell for cell in row if normalize_text(cell)]
        if non_empty and sum(isinstance(cell, str) for cell in non_empty) >= max(1, len(non_empty) // 2):
            return index
    return 1


def _redacted(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (int, float, bool)):
        return value
    return redact_sample(normalize_text(value))


def _fingerprint(sheet: str, headers: list[str], values: Iterable[Any]) -> str:
    payload = {
        "sheet": sheet,
        "headers": headers,
        "values": [normalize_text(value) for value in values],
    }
    return hashlib.sha256(json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode()).hexdigest()


def inspect_bytes(source: bytes, *, max_rows: int | None = None, sample_limit: int | None = None) -> dict[str, Any]:
    """Inspect an XLSX byte string and return a bounded JSON-serializable manifest."""

    max_rows = max_rows or settings.max_data_rows
    sample_limit = sample_limit or settings.max_samples_per_column
    workbook = load_workbook(BytesIO(source), read_only=True, data_only=False)
    try:
        sheets: list[dict[str, Any]] = []
        total_rows = 0
        for worksheet in workbook.worksheets:
            rows = list(worksheet.iter_rows(values_only=True))
            if len(rows) > max_rows:
                raise ValueError(f"workbook exceeds the {max_rows} row limit")
            header_index = _header_row(rows)
            raw_headers = list(rows[header_index - 1]) if rows else []
            width = max(len(raw_headers), worksheet.max_column or 0)
            headers = [normalize_text(raw_headers[i]) or f"Column_{i + 1}" for i in range(width)]
            data_rows = [tuple(row[i] if i < len(row) else None for i in range(width)) for row in rows[header_index:]]
            column_info = []
            formula_locations: list[str] = []
            blank_rows: list[int] = []
            fingerprints: list[dict[str, Any]] = []
            for row_number, row in enumerate(data_rows, start=header_index + 1):
                if all(normalize_text(cell) == "" for cell in row):
                    blank_rows.append(row_number)
                for column_number, cell in enumerate(row, start=1):
                    if isinstance(cell, str) and cell.startswith("="):
                        formula_locations.append(f"{get_column_letter(column_number)}{row_number}")
                fingerprints.append({"row_number": row_number, "fingerprint": _fingerprint(worksheet.title, headers, row)})
            for column_number, header in enumerate(headers):
                values = [row[column_number] for row in data_rows]
                distribution = Counter(_type_name(value) for value in values if normalize_text(value) != "")
                samples = [_redacted(value) for value in values if normalize_text(value) != ""][:sample_limit]
                column_info.append({
                    "index": column_number + 1,
                    "letter": get_column_letter(column_number + 1),
                    "name": header,
                    "normalized_name": normalize_text(header).casefold(),
                    "type_distribution": dict(distribution),
                    "blank_count": sum(normalize_text(value) == "" for value in values),
                    "samples": samples,
                })
            blank_columns = [
                column["name"] for column in column_info
                if column["blank_count"] == len(data_rows) and data_rows
            ]
            sheets.append({
                "name": worksheet.title,
                "max_row": worksheet.max_row or 0,
                "max_column": worksheet.max_column or 0,
                "header_row": header_index,
                "headers": headers,
                "columns": column_info,
                "row_fingerprints": fingerprints,
                "blank_row_numbers": blank_rows,
                "blank_columns": blank_columns,
                "formula_locations": formula_locations,
            })
            total_rows += len(data_rows)
        return {
            "manifest_version": "1",
            "format": "xlsx",
            "sheet_count": len(sheets),
            "total_rows": total_rows,
            "sheets": sheets,
        }
    finally:
        workbook.close()


def manifest_json(manifest: dict[str, Any]) -> bytes:
    return json.dumps(manifest, indent=2, sort_keys=True).encode("utf-8")
