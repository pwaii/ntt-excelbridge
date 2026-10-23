from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.workers.apply_changes import apply_accepted_updates, apply_workbook_changes
from app.workers.formulas import verify_formula
from app.workers.inspect_workbook import inspect_bytes
from app.workers.matching import build_tool5_inputs, exact_match_rows, generate_candidates
from app.workers.quality_rules import run_quality_rules
from scripts.make_fixtures import write_fixtures


def _fixture_bytes(tmp_path):
    current, main = write_fixtures(tmp_path)
    return current.read_bytes(), main.read_bytes()


def test_inspection_manifest_contains_headers_types_fingerprints_and_formulas(tmp_path):
    source, _ = _fixture_bytes(tmp_path)
    manifest = inspect_bytes(source)
    sheet = manifest["sheets"][0]
    assert sheet["name"] == "Customers"
    assert "CustomerID" in sheet["headers"]
    assert sheet["row_fingerprints"]
    assert sheet["columns"][0]["type_distribution"]
    assert "Notes" in sheet["blank_columns"]


def test_quality_rules_detect_duplicate_blank_and_type_issues(tmp_path):
    source, _ = _fixture_bytes(tmp_path)
    report = run_quality_rules(source)
    kinds = {issue["kind"] for issue in report["issues"]}
    assert {"duplicate_rows", "blank_row", "blank_column", "type_mismatch"} <= kinds


def test_formula_verifier_allows_safe_formula_and_rejects_external_or_volatile():
    result = verify_formula("=IF(A2>0,B2,0)", min_row=2, max_row=10, max_column=3, target_column=3)
    assert result["status"] == "syntax_verified"
    with pytest.raises(ValueError):
        verify_formula("=HYPERLINK(A2)", min_row=2, max_row=10, max_column=3)
    with pytest.raises(ValueError):
        verify_formula("=[other.xlsx]Sheet1!A2", min_row=2, max_row=10, max_column=3)


def test_exact_match_returns_ambiguous_and_unmatched_without_inventing_ids():
    current = [{"row_id": "C1", "CustomerID": "000100"}, {"row_id": "C2", "CustomerID": "999999"}]
    main = [{"row_id": "M1", "CustomerID": "000100"}, {"row_id": "M2", "CustomerID": "000100"}]
    results = exact_match_rows(current, main, "CustomerID")
    assert results[0]["match_type"] == "ambiguous"
    assert results[0]["selected_main_row_id"] is None
    assert results[1]["match_type"] == "unmatched"


def test_apply_writes_new_bytes_and_preserves_formula_origin():
    from openpyxl import Workbook

    source_workbook = Workbook()
    sheet = source_workbook.active
    sheet.title = "Customers"
    sheet.append(["Quantity", "UnitPrice", "Total"])
    sheet.append([2, 10, None])
    source_stream = BytesIO()
    source_workbook.save(source_stream)
    source = source_stream.getvalue()
    output, audit = apply_workbook_changes(source, formula={"formula_template": "=A{row}*B{row}", "target_header": "Total", "target_column": 3, "start_row": 2, "end_row": 2})
    assert output != source
    reopened = load_workbook(BytesIO(output), data_only=False)
    assert reopened.active["C2"].value == "=A2*B2"
    assert any(item["origin"] == "approved-formula" for item in audit)


def test_candidate_generation_is_bounded_and_keeps_row_ids():
    current = [{"row_id": "Current!2", "CustomerID": "ACME-100"}]
    main = [
        {"row_id": "Main!2", "CustomerID": "ACME-100"},
        {"row_id": "Main!3", "CustomerID": "ACME-101"},
        {"row_id": "Main!4", "CustomerID": "ACME-102"},
    ]
    result = generate_candidates(current, main, "CustomerID", max_candidates=2)
    assert result[0]["match_type"] == "exact"
    assert result[0]["selected_main_row_id"] == "Main!2"
    inputs = build_tool5_inputs(current, main, "CustomerID", max_candidates=2)
    assert inputs["deterministic_matches"][0]["current_row_id"] == "Current!2"
    assert all(item.startswith("Main!") for item in inputs["deterministic_matches"][0]["candidate_main_row_ids"])


def test_tool4_inputs_are_limited_to_shared_fields(tmp_path):
    current_path, main_path = write_fixtures(tmp_path)

    class FixtureS3:
        def get_bytes(self, key):
            return current_path.read_bytes() if key == "current" else main_path.read_bytes()

    from app.lambda_handlers import _tool_inputs

    inputs = _tool_inputs(
        {
            "config": {"input_sheet": "Customers", "main_sheet": "CanonicalCustomers"},
            "files": {"input": {"key": "current"}, "main": {"key": "main"}},
        },
        FixtureS3(),
        "tool4",
    )
    fields = {item["field"] for item in inputs["source_values"]}
    assert fields == {"CustomerID", "FullName", "Country"}
    assert {item["value"] for item in inputs["source_values"]} >= {"United States", "U.S.A."}
    assert all(column["name"] in fields for column in inputs["samples"])


def test_apply_accepted_updates_copies_values_and_skips_duplicate_destinations():
    from openpyxl import Workbook

    current_book = Workbook()
    current_sheet = current_book.active
    current_sheet.title = "Current"
    current_sheet.append(["CustomerID", "Status"])
    current_sheet.append(["ACME-100", "Active"])
    current_stream = BytesIO()
    current_book.save(current_stream)

    main_book = Workbook()
    main_sheet = main_book.active
    main_sheet.title = "Main"
    main_sheet.append(["CustomerID", "Status"])
    main_sheet.append(["ACME-100", "Old"])
    main_stream = BytesIO()
    main_book.save(main_stream)

    matches = [{
        "decision_id": "tool5-hash:item:0",
        "current_row_id": "Current!2",
        "selected_main_row_id": "Main!2",
        "candidate_main_row_ids": ["Main!2"],
        "match_type": "exact",
    }]
    output, audit = apply_accepted_updates(
        main_stream.getvalue(),
        current_stream.getvalue(),
        matches,
        {"tool5-hash:item:0": {"decision": "accept"}},
        key_column="CustomerID",
        main_sheet="Main",
        current_sheet="Current",
    )
    reopened = load_workbook(BytesIO(output), data_only=False)
    assert reopened["Main"]["B2"].value == "Active"
    assert any(item["action"] == "updated" for item in audit)


def test_apply_accepted_updates_prior_destinations_blocks_cross_file_overwrite():
    from openpyxl import Workbook

    def _book(title, rows):
        book = Workbook()
        ws = book.active
        ws.title = title
        for row in rows:
            ws.append(row)
        stream = BytesIO()
        book.save(stream)
        return stream.getvalue()

    main = _book("Main", [["CustomerID", "Status"], ["A-1", "Old"]])
    file_a = _book("Current", [["CustomerID", "Status"], ["A-1", "FromA"]])
    file_b = _book("Current", [["CustomerID", "Status"], ["A-1", "FromB"]])

    def match(did):
        return [{"decision_id": did, "current_row_id": "Current!2", "selected_main_row_id": "Main!2"}]

    def accept(did):
        return {did: {"decision": "accept"}}

    # Sequential merge of two files into one main, sharing the destination set.
    seen: set = set()
    merged, _ = apply_accepted_updates(
        main, file_a, match("a"), accept("a"),
        key_column="CustomerID", main_sheet="Main", current_sheet="Current",
        prior_destinations=seen,
    )
    merged, audit_b = apply_accepted_updates(
        merged, file_b, match("b"), accept("b"),
        key_column="CustomerID", main_sheet="Main", current_sheet="Current",
        prior_destinations=seen,
    )

    # File A won; File B's cross-file collision was skipped, not silently applied.
    assert load_workbook(BytesIO(merged), data_only=False)["Main"]["B2"].value == "FromA"
    assert any(item["action"] == "skipped" for item in audit_b)


def test_apply_accepted_updates_requires_explicit_conflict_policy_behavior():
    from openpyxl import Workbook

    current_book = Workbook()
    current_sheet = current_book.active
    current_sheet.title = "Current"
    current_sheet.append(["CustomerID", "Status"])
    current_sheet.append(["A-1", "First"])
    current_sheet.append(["A-1", "Last"])
    current_stream = BytesIO()
    current_book.save(current_stream)

    main_book = Workbook()
    main_sheet = main_book.active
    main_sheet.title = "Main"
    main_sheet.append(["CustomerID", "Status"])
    main_sheet.append(["A-1", "Old"])
    main_stream = BytesIO()
    main_book.save(main_stream)

    matches = [
        {"decision_id": "p:0", "current_row_id": "Current!2", "selected_main_row_id": "Main!2"},
        {"decision_id": "p:1", "current_row_id": "Current!3", "selected_main_row_id": "Main!2"},
    ]
    decisions = {"p:0": {"decision": "accept"}, "p:1": {"decision": "accept"}}
    rejected, rejected_audit = apply_accepted_updates(
        main_stream.getvalue(), current_stream.getvalue(), matches, decisions,
        key_column="CustomerID", main_sheet="Main", current_sheet="Current",
    )
    assert load_workbook(BytesIO(rejected), data_only=False)["Main"]["B2"].value == "Old"
    assert sum(item["reason"] == "conflict policy reject_all" for item in rejected_audit) == 2

    latest, _ = apply_accepted_updates(
        main_stream.getvalue(), current_stream.getvalue(), matches, decisions,
        key_column="CustomerID", main_sheet="Main", current_sheet="Current",
        conflict_policy="latest_wins",
    )
    assert load_workbook(BytesIO(latest), data_only=False)["Main"]["B2"].value == "Last"
