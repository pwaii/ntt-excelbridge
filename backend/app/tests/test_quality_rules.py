"""Deterministic quality-rule tests."""

from io import BytesIO

from openpyxl import Workbook
import pytest

from app.workers.quality_rules import run_quality_rules


def _workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Customers"
    sheet.append(["CustomerID", "Amount", "Unused"])
    sheet.append(["A-1", 10, None])
    sheet.append(["A-1", "not-a-number", None])
    sheet.append([None, None, None])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def test_quality_rules_are_stable_and_include_recommended_actions():
    first = run_quality_rules(_workbook_bytes())
    second = run_quality_rules(_workbook_bytes())
    assert first == second
    assert first["issue_count"] >= 3
    assert all(issue["recommended_action"] for issue in first["issues"])


def _single_column(values: list) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet"
    sheet.append(["Qty"])
    for value in values:
        sheet.append([value])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def test_single_poison_value_is_flagged_despite_strong_majority():
    # One text value in 100 integers: the old 20%-minority gate suppressed this (AI-001).
    report = run_quality_rules(_single_column([*([1] * 100), "two"]))
    issue = next(issue for issue in report["issues"] if issue["kind"] == "type_mismatch")
    assert issue["column"] == "Qty"
    assert issue["proposed_target_type"] == "integer"


def test_numeric_column_with_decimals_is_not_a_type_mismatch():
    # int + decimal are the same numeric group — a price column must not be flagged (AI-001).
    report = run_quality_rules(_single_column([10, 10.5, 12, 9.99]))
    assert not [issue for issue in report["issues"] if issue["kind"] == "type_mismatch"]


def test_blank_key_field_flagged_on_otherwise_populated_row():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Parts"
    sheet.append(["PartNo", "Description", "OnHand"])
    sheet.append(["P-1", "Bolt", 10])
    sheet.append([None, "Nut", 5])  # blank key, other fields populated
    stream = BytesIO()
    workbook.save(stream)
    report = run_quality_rules(stream.getvalue())
    blank_key = next(issue for issue in report["issues"] if issue["kind"] == "blank_key_field")
    assert blank_key["row_numbers"] == [3]
    assert blank_key["column"] == "PartNo"


@pytest.mark.parametrize("value,expected", [(10, "integer"), (True, "boolean"), ("A123", "string")])
def test_mixed_types_use_canonical_target_names(value, expected):
    workbook = Workbook()
    workbook.active.append(["Value"])
    workbook.active.append([value])
    workbook.active.append(["unparseable"])
    stream = BytesIO()
    workbook.save(stream)
    report = run_quality_rules(stream.getvalue())
    issue = next(issue for issue in report["issues"] if issue["kind"] == "type_mismatch")
    assert issue["proposed_target_type"] == expected
