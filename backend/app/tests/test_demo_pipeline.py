"""Tests for the AWS-backed pass-through output pipeline."""

from __future__ import annotations

from copy import deepcopy
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.domain.enums import JobStatus
from app.storage.dynamo_store import RevisionConflictError
from app.workers.demo_pipeline import run_demo_pipeline


def _workbook_bytes() -> bytes:
    workbook = Workbook()
    workbook.active["A1"] = "Status"
    workbook.active["A2"] = "Ready"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class FakeDynamo:
    def __init__(self):
        self.item = {
            "job_id": "job-1",
            "owner_id": "local-user",
            "status": JobStatus.INSPECTING.value,
            "revision": 2,
            "config": {"comment_name": "Status"},
            "files": {
                "input": {
                    "filename": "input.xlsx",
                    "key": "jobs/job-1/input/input/input.xlsx",
                    "uploaded": True,
                }
            },
            "created_at": "2026-10-15T00:00:00+00:00",
            "updated_at": "2026-10-15T00:00:00+00:00",
            "expires_at": 1_800_000_000,
        }

    def get_job(self, job_id):
        return deepcopy(self.item) if job_id == "job-1" else None

    def update_job(self, job_id, expected_revision, updates):
        if expected_revision != self.item["revision"]:
            raise RevisionConflictError("stale")
        self.item.update(deepcopy(updates))
        self.item["revision"] += 1
        return deepcopy(self.item)


class FakeS3:
    def __init__(self):
        self.objects = {"jobs/job-1/input/input/input.xlsx": _workbook_bytes()}

    def get_bytes(self, key):
        return self.objects[key]

    def put_bytes(self, key, body, content_type):
        self.objects[key] = body

    def artifact_key(self, job_id, area, filename):
        return f"jobs/{job_id}/{area}/{filename}"


def test_demo_pipeline_writes_and_validates_output():
    dynamo = FakeDynamo()
    s3 = FakeS3()

    completed = run_demo_pipeline("job-1", 2, dynamo, s3)

    assert completed["status"] == JobStatus.COMPLETED.value
    assert completed["artifacts"]["output"]["key"] == "jobs/job-1/output/final.xlsx"
    output = s3.objects["jobs/job-1/output/final.xlsx"]
    workbook = load_workbook(BytesIO(output), read_only=True, data_only=False)
    assert workbook.active["A1"].value == "Status"
    workbook.close()
    assert "jobs/job-1/inspection/manifest.json" in s3.objects
