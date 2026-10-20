"""Self-check: generated varied-domain samples carry their labeled problems.
Tests the fixtures (the deliverable), not generator internals — generate into a
tmp dir, reopen, assert each condition. See
docs/superpowers/specs/2026-10-20-varied-sample-fixtures-design.md."""
from datetime import datetime

import pytest
from openpyxl import load_workbook

from scripts.make_samples import write_samples


@pytest.fixture
def books(tmp_path):
    """filename -> loaded Workbook for every generated sample."""
    return {p.name: load_workbook(p) for p in write_samples(tmp_path)}


def _rows(ws):
    """Row tuples below the header, values only."""
    return list(ws.iter_rows(min_row=2, values_only=True))


def _formula_text_cells(ws):
    """Cells whose stored value starts with '=' (injection / formula-as-data)."""
    return [c for row in ws.iter_rows(min_row=2) for c in row
            if isinstance(c.value, str) and c.value.startswith("=")]


# --- finance.xlsx / Invoices ---

def test_finance_duplicate_and_blank_row(books):
    rows = _rows(books["finance.xlsx"]["Invoices"])
    assert rows[0] == rows[1]                                      # F1 == F2
    assert sum(1 for r in rows if all(c is None for c in r)) == 1  # one blank row


def test_finance_memo_is_blank_column(books):
    rows = _rows(books["finance.xlsx"]["Invoices"])
    assert all(r[7] is None for r in rows)                         # Memo (index 7)


def test_finance_unitprice_type_mismatch(books):
    prices = [r[4] for r in _rows(books["finance.xlsx"]["Invoices"]) if r[4] is not None]
    assert any(isinstance(p, float) for p in prices)
    assert any(isinstance(p, str) for p in prices)


def test_finance_injection_cell_stored_as_text(books):
    cells = _formula_text_cells(books["finance.xlsx"]["Invoices"])
    assert cells and all(c.data_type == "s" for c in cells)


# --- hr.xlsx / Employees ---

def test_hr_duplicate_and_blank_row(books):
    rows = _rows(books["hr.xlsx"]["Employees"])
    assert rows[0] == rows[1]                                      # H1 == H2
    assert sum(1 for r in rows if all(c is None for c in r)) == 1  # one blank row


def test_hr_salary_type_mismatch(books):
    salaries = [r[5] for r in _rows(books["hr.xlsx"]["Employees"]) if r[5] is not None]
    assert any(isinstance(s, int) for s in salaries)
    assert "N/A" in salaries


def test_hr_fulltime_boolean_like(books):
    ft = {r[6] for r in _rows(books["hr.xlsx"]["Employees"]) if r[6] is not None}
    assert {"TRUE", "Yes", 1, "N"}.issubset(ft)


def test_hr_hiredate_has_invalid_value(books):
    dates = [r[4] for r in _rows(books["hr.xlsx"]["Employees"]) if r[4] is not None]

    def parseable(s):
        for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
            try:
                datetime.strptime(s, fmt)
                return True
            except ValueError:
                pass
        return False

    assert any(not parseable(d) for d in dates)                    # "2023-13-45"


def test_hr_unicode_and_nonlatin_names(books):
    firsts = {r[1] for r in _rows(books["hr.xlsx"]["Employees"]) if r[1]}
    assert "José" in firsts                                        # accented
    assert "李" in firsts                                          # non-latin
