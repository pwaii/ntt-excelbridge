"""Transformation and output-validation tests."""

from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.workers.apply_changes import apply_accepted_changes
from app.workers.validate_output import validate_output


def _source() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Customers"
    sheet.append(["CustomerID", "Status"])
    sheet.append(["A-1", "Active"])
    sheet.append(["A-1", "Active"])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def test_accepted_duplicate_deletion_creates_valid_new_output():
    source = _source()
    output, audit = apply_accepted_changes(
        source,
        quality_report={
            "issues": [{
                "issue_id": "duplicate-1",
                "kind": "duplicate_rows",
                "row_numbers": [2, 3],
            }]
        },
        decisions={"duplicate-1": {"decision": "accept"}},
    )
    assert output != source
    workbook = load_workbook(BytesIO(output), data_only=False)
    assert workbook.active.max_row == 2
    assert any(item.get("action") == "delete" for item in audit)
    validation = validate_output(source, output, expected_sheet="Customers", expected_headers=["CustomerID", "Status"])
    assert validation["status"] == "valid"
    assert validation["source_and_output_equal"] is False
