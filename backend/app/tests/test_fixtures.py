"""Self-check: the generated §16.3 fixtures actually carry all 12 conditions.
Tests the fixture (the deliverable), not generator internals — generate into a tmp
dir, reopen, and assert each condition. See
docs/superpowers/specs/2026-10-15-ai-day2-fixtures-design.md."""
from datetime import datetime

import pytest
from openpyxl import load_workbook

from scripts.make_fixtures import write_fixtures


@pytest.fixture
def workbooks(tmp_path):
    current_path, main_path = write_fixtures(tmp_path)
    cur = load_workbook(current_path)["Customers"]
    main = load_workbook(main_path)["CanonicalCustomers"]
    return cur, main


def _data_rows(ws):
    """Row tuples below the header, values only."""
    return list(ws.iter_rows(min_row=2, values_only=True))


def test_customer_id_leading_zeros_preserved(workbooks):
    cur, _ = workbooks
    first = _data_rows(cur)[0]
    assert first[0] == "000100"
    assert isinstance(first[0], str)


def test_notes_column_is_blank(workbooks):
    cur, _ = workbooks
    assert all(row[7] is None for row in _data_rows(cur))


def test_exactly_one_completely_blank_row(workbooks):
    cur, _ = workbooks
    blanks = [r for r in _data_rows(cur) if all(c is None for c in r)]
    assert len(blanks) == 1


def test_two_exact_duplicate_rows(workbooks):
    cur, _ = workbooks
    rows = _data_rows(cur)
    assert rows[0] == rows[1]


def test_date_column_has_two_formats_and_one_invalid(workbooks):
    cur, _ = workbooks
    dates = [r[3] for r in _data_rows(cur) if r[3] is not None]
    assert "2023-01-15" in dates      # ISO variant
    assert "15/02/2023" in dates      # DD/MM/YYYY variant

    def parseable(s):
        for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
            try:
                datetime.strptime(s, fmt)
                return True
            except ValueError:
                pass
        return False

    assert any(not parseable(d) for d in dates)   # "2023-13-45" is invalid


def test_boolean_like_column_variants(workbooks):
    cur, _ = workbooks
    rows = _data_rows(cur)
    actives = {r[4] for r in rows if r[4] is not None}
    assert {"Y", "Yes", "N", "No"}.issubset(actives)
    # a populated row whose Active is blank (not the fully-blank row)
    assert any(r[4] is None and r[1] is not None for r in rows)


def test_ambiguous_and_shared_main_row_pairs(workbooks):
    cur, main = workbooks
    main_rows = _data_rows(main)
    cur_rows = _data_rows(cur)
    # ambiguity source: two identical names in main
    assert sum(1 for r in main_rows if r[1] == "Jane Smith") == 2
    # shared-main-row pair: two current rows, one main row, same key
    assert sum(1 for r in cur_rows if r[0] == "000500") == 2
    assert sum(1 for r in main_rows if r[0] == "000500") == 1


def test_unmapped_standardization_value(workbooks):
    cur, main = workbooks
    main_countries = {r[2] for r in _data_rows(main)}
    cur_countries = {r[2] for r in _data_rows(cur) if r[2] is not None}
    assert "Wakanda" in cur_countries
    assert "Wakanda" not in main_countries
