"""Actual API/worker/workbook integration; AWS and Bedrock boundaries are synthetic."""

import json
from copy import deepcopy
from io import BytesIO

import pytest
from botocore.exceptions import ClientError
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app import lambda_handlers as handlers
from app.api.dependencies import get_dynamo_store, get_s3_store, get_workflow_store
from app.auth import current_user_id
from app.main import app
from app.tests.test_api import FakeDynamo
from app.tests.test_workbook_configuration import WorkbookS3


class Store(FakeDynamo):
    def __init__(self):
        super().__init__()
        self.children = {}

    def put_child(self, job_id, sort_key, item):
        self.children[sort_key] = deepcopy(item)
        return deepcopy(item)

    def list_children_page(self, job_id, prefix, limit=100, cursor=None):
        rows = [
            deepcopy(value)
            for key, value in sorted(self.children.items())
            if key.startswith(prefix)
        ]
        start = int(cursor or 0)
        end = start + min(limit, 2)  # Exercise every producer/consumer over multiple pages.
        return rows[start:end], str(end) if end < len(rows) else None


class Storage(WorkbookS3):
    def object_exists(self, key):
        return key in self.objects


class Workflow:
    state_machine_arn = "arn:synthetic:state-machine"

    def __init__(self):
        self.started = []

    def start_segment(self, job_id, segment, revision, **kwargs):
        self.started.append(segment)
        return {"execution_arn": f"arn:synthetic:execution:{segment}"}


def workbook(main=False):
    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append(["Synthetic report"])
    sheet.append([])
    sheet.append(["ID", "Amount", "Status"])
    sheet.append(["001", 0 if main else 2, "Active" if main else "Old"])
    if not main:
        sheet.append(["999", 3, "Old"])
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


@pytest.fixture
def api(monkeypatch):
    store, storage, workflow = Store(), Storage(), Workflow()
    app.dependency_overrides[get_dynamo_store] = lambda: store
    app.dependency_overrides[get_s3_store] = lambda: storage
    app.dependency_overrides[get_workflow_store] = lambda: workflow
    app.dependency_overrides[current_user_id] = lambda: "local-user"
    monkeypatch.setattr(handlers, "_stores", lambda: (store, storage))
    with TestClient(app) as client:
        yield client, store, storage, workflow
    app.dependency_overrides.clear()


def proposal(store, storage, tool, payload, index=None):
    identity = f"synthetic-{tool}" + (f":item:{index}" if index is not None else "")
    key = storage.artifact_key(store.item["job_id"], "proposals", f"{tool}.json")
    storage.put_bytes(key, json.dumps({"payload": payload}).encode(), "application/json")
    record = {"proposal_id": identity, "tool": tool, "artifact_key": key}
    if index is not None:
        record["proposal_index"] = index
    store.put_child(store.item["job_id"], f"PROPOSAL#{identity}", record)
    return identity


def report(client, store, stage, limit=100):
    rows, cursor = [], 0
    while True:
        response = client.get(
            f"/v1/jobs/{store.item['job_id']}/review-report",
            params={
                "stage": stage,
                "expected_revision": store.item["revision"],
                "cursor": cursor,
                "limit": limit,
            },
        )
        assert response.status_code == 200, response.text
        page = response.json()
        rows.extend(page["items"])
        if "nextCursor" not in page:
            assert page["complete"] and len(rows) == page["total"]
            return rows
        cursor = page["nextCursor"]


def prepare(api, standardize=False, update=False):
    client, store, storage, _workflow = api
    body = {
        "config": {
            "comment_name": "Double Amount",
            "output_header": "Double",
            "standardize": standardize,
            "update": update,
        },
        "input_filename": "current.xlsx",
    }
    if standardize or update:
        body["main_filename"] = "main.xlsx"
    response = client.post("/v1/jobs", json=body)
    assert response.status_code == 201, response.text
    job_id = response.json()["job"]["job_id"]
    selections = {}
    for slot, file in store.item["files"].items():
        storage.objects[file["key"]] = workbook(slot == "main")
        assert (
            client.post(f"/v1/jobs/{job_id}/uploads/complete", json={"slot": slot}).status_code
            == 200
        )
        metadata_response = client.get(
            f"/v1/jobs/{job_id}/workbooks/{slot}/metadata",
            params={"sheet_name": "Data", "header_row": 3},
        )
        assert metadata_response.status_code == 200, metadata_response.text
        metadata = metadata_response.json()
        selections[slot] = {
            "sheet_name": "Data",
            "header_row": 3,
            "header_confirmed": True,
            "key_column": "A" if update else None,
            "workbook_version": metadata["workbook_version"],
        }
    assert (
        client.put(
            f"/v1/jobs/{job_id}/configuration",
            json={"expected_revision": store.item["revision"], "selections": selections},
        ).status_code
        == 200
    )
    if update:
        assert store.item["config"]["key_column"] == "ID"
    return job_id


@pytest.mark.parametrize(
    "standardize,update", [(False, False), (True, False), (False, True), (True, True)]
)
def test_upload_configuration_both_review_gates_and_download_all_modes(api, standardize, update):
    client, store, storage, workflow = api
    job_id = prepare(api, standardize, update)
    original = dict(storage.objects)
    assert (
        client.post(
            f"/v1/jobs/{job_id}/start", params={"expected_revision": store.item["revision"]}
        ).status_code
        == 200
    )
    handlers.inspect_handler({"job_id": job_id}, None)
    handlers.quality_handler({"job_id": job_id}, None)
    proposal(
        store,
        storage,
        "tool1",
        {
            "meaning": "Double Amount",
            "interpretation": "Multiply Amount by two",
            "confidence": 0.99,
        },
    )
    handlers.quality_gate_handler({"job_id": job_id}, None)
    findings = report(client, store, "quality", limit=1)
    decisions = {row["id"]: {"decision": "accept"} for row in findings if row["requiresDecision"]}
    revision = store.item["revision"]
    response = client.post(
        f"/v1/jobs/{job_id}/reviews", json={"expected_revision": revision, "decisions": decisions}
    )
    assert response.status_code == 200, response.text
    assert workflow.started == [1, 2]
    assert (
        client.post(
            f"/v1/jobs/{job_id}/reviews",
            json={"expected_revision": revision, "decisions": decisions},
        ).status_code
        == 409
    )
    formula = proposal(
        store,
        storage,
        "tool3",
        {
            "formula_template": "=B{row}*2",
            "target_header": "Double",
            "referenced_headers": ["Amount"],
            "explanation": "Double the amount",
            "confidence": 0.99,
        },
    )
    if standardize:
        proposal(
            store,
            storage,
            "tool4",
            {
                "mappings": [
                    {
                        "source_field": "Status",
                        "source_value": "Old",
                        "canonical_field": "Status",
                        "canonical_value": "Active",
                        "reason": "Use main vocabulary",
                        "confidence": 0.99,
                    }
                ]
            },
            0,
        )
    if update:
        matches = [
            {
                "current_row_id": "Data!4",
                "candidate_main_row_ids": ["Data!4"],
                "selected_main_row_id": "Data!4",
                "match_type": "exact",
                "reason": "Same ID",
                "confidence": 1,
            },
            {
                "current_row_id": "Data!5",
                "candidate_main_row_ids": [],
                "selected_main_row_id": None,
                "match_type": "unmatched",
                "reason": "No target; skip",
                "confidence": 1,
            },
        ]
        for index in range(len(matches)):
            proposal(store, storage, "tool5", {"matches": matches}, index)
    handlers.quality_gate_handler({"job_id": job_id, "gate": "transformation"}, None)
    proposals = report(client, store, "transformation", limit=1)
    formula_row = next(row for row in proposals if row["id"] == formula)
    assert formula_row["formula"]["syntax"] == "verified"
    decisions = {row["id"]: {"decision": "accept"} for row in proposals if row["requiresDecision"]}
    revision = store.item["revision"]
    assert (
        client.post(
            f"/v1/jobs/{job_id}/reviews",
            json={"expected_revision": revision, "decisions": decisions},
        ).status_code
        == 200
    )
    assert workflow.started == [1, 2, 3]
    handlers.apply_handler({"job_id": job_id}, None)
    handlers.validate_handler({"job_id": job_id}, None)
    result = client.get(f"/v1/jobs/{job_id}").json()["job"]
    assert result["status"] == "COMPLETED"
    assert result["result_summary"]["updated_rows"] == (1 if update else 0)
    assert all(storage.objects[key] == value for key, value in original.items())
    final = load_workbook(
        BytesIO(storage.objects[result["artifacts"]["output"]["key"]]), data_only=False
    )
    assert final["Data"].max_row == (4 if update else 5)  # Unmatched rows are never inserted.
    assert final["Data"]["B4"].value == 2
    assert final["Data"]["D4"].value == "=B4*2"
    if standardize:
        assert final["Data"]["C4"].value == "Active"
    final.close()
    for artifact_type in ["output", "audit", "review_summary", "validation"]:
        assert client.get(f"/v1/jobs/{job_id}/artifacts/{artifact_type}").status_code == 200
    persisted = report(client, store, "transformation")
    assert all(row["status"] != "proposed" for row in persisted if row["id"] in decisions)


def _dup_workbook():
    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append(["ID", "Amount", "Status"])
    sheet.append(["001", 2, "Old"])
    sheet.append(["001", 2, "Old"])  # exact duplicate -> a quality issue every slot will raise
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def test_fanout_handlers_isolate_each_input_slot(api):
    client, store, storage, _ = api
    response = client.post(
        "/v1/jobs",
        json={
            "config": {
                "comment_name": "Double Amount",
                "output_header": "Double",
                "standardize": False,
                "update": False,
            },
            "input_filenames": ["current.xlsx", "current-2.xlsx"],
        },
    )
    assert response.status_code == 201, response.text
    job_id = response.json()["job"]["job_id"]
    assert set(store.item["files"]) == {"input", "input_2"}
    for slot, file in store.item["files"].items():
        storage.objects[file["key"]] = _dup_workbook()
        assert (
            client.post(f"/v1/jobs/{job_id}/uploads/complete", json={"slot": slot}).status_code
            == 200
        )
    store.item["config"].update({"input_sheet": "Data", "input_header_row": 1})

    # The Map fan-out: inspect + quality run once per input slot.
    for slot in ("input", "input_2"):
        handlers.inspect_handler({"job_id": job_id, "slot": slot}, None)
        handlers.quality_handler({"job_id": job_id, "slot": slot}, None)

    artifacts = store.item["artifacts"]
    # Each slot wrote its own isolated inspection + quality artifacts.
    assert artifacts["inspection"]["key"] != artifacts["inspection_input_2"]["key"]
    assert artifacts["quality"]["key"] != artifacts["quality_input_2"]["key"]
    # Issues are namespaced by slot; both slots produced some and record their slot.
    assert any(sk.startswith("ISSUE#input_2#") for sk in store.children)
    assert any(
        sk.startswith("ISSUE#") and not sk.startswith("ISSUE#input_2#") for sk in store.children
    )
    assert {child["slot"] for child in store.children.values() if child.get("entity_type") == "issue"} == {
        "input",
        "input_2",
    }


def test_review_rejects_unknown_incomplete_and_stale_decisions(api):
    client, store, storage, _ = api
    job_id = prepare(api)
    proposal(
        store,
        storage,
        "tool1",
        {"meaning": "Meaning", "interpretation": "Inspect the meaning", "confidence": 1},
    )
    handlers.quality_handler({"job_id": job_id}, None)
    handlers.quality_gate_handler({"job_id": job_id}, None)
    revision = store.item["revision"]
    endpoint = f"/v1/jobs/{job_id}/reviews"
    assert (
        client.post(
            endpoint,
            json={"expected_revision": revision, "decisions": {"unknown": {"decision": "accept"}}},
        ).status_code
        == 422
    )
    assert store.item["revision"] == revision
    assert (
        client.get(
            f"/v1/jobs/{job_id}/review-report",
            params={"stage": "quality", "expected_revision": revision - 1},
        ).status_code
        == 409
    )


def ready_quality(api):
    _client, store, storage, _workflow = api
    job_id = prepare(api)
    identity = proposal(
        store,
        storage,
        "tool1",
        {"meaning": "Meaning", "interpretation": "Inspect the meaning", "confidence": 1},
    )
    handlers.quality_handler({"job_id": job_id}, None)
    handlers.quality_gate_handler({"job_id": job_id}, None)
    return job_id, identity


def test_gate_requires_complete_decisions_and_ignore_reasons(api):
    client, store, storage, _ = api
    job_id, identity = ready_quality(api)
    rules = [
        {
            "pattern": "Numeric text",
            "example_original": "2",
            "normalized_form": "2",
            "reason": "Numeric format",
            "confidence": 1,
        }
    ]
    proposal(store, storage, "tool2", {"column_id": "Amount", "rules": rules}, 0)
    revision = store.item["revision"]
    endpoint = f"/v1/jobs/{job_id}/reviews"
    assert (
        client.post(
            endpoint,
            json={"expected_revision": revision, "decisions": {identity: {"decision": "accept"}}},
        ).status_code
        == 422
    )
    rows = report(client, store, "quality")
    decisions = {
        row["id"]: {"decision": "ignore", "reason": "  "} for row in rows if row["requiresDecision"]
    }
    assert (
        client.post(
            endpoint, json={"expected_revision": revision, "decisions": decisions}
        ).status_code
        == 422
    )
    assert store.item["revision"] == revision


def test_report_ownership_is_checked_before_loading_artifacts(api):
    client, store, _, _ = api
    job_id, _ = ready_quality(api)
    app.dependency_overrides[current_user_id] = lambda: "another-owner"
    assert (
        client.get(
            f"/v1/jobs/{job_id}/review-report",
            params={"stage": "quality", "expected_revision": store.item["revision"]},
        ).status_code
        == 404
    )


def test_review_persists_even_if_the_secondary_index_write_fails(api, monkeypatch):
    client, store, _, _ = api
    job_id, identity = ready_quality(api)
    monkeypatch.setattr(
        store,
        "put_child",
        lambda *args: (_ for _ in ()).throw(RuntimeError("synthetic index outage")),
    )
    response = client.post(
        f"/v1/jobs/{job_id}/reviews",
        json={
            "expected_revision": store.item["revision"],
            "decisions": {identity: {"decision": "accept"}},
        },
    )
    assert response.status_code == 200
    persisted = report(client, store, "quality")
    assert next(row for row in persisted if row["id"] == identity)["status"] == "human_accepted"


def test_failed_continuation_is_recoverable_without_replaying_review(api, monkeypatch):
    client, store, _, workflow = api
    job_id, identity = ready_quality(api)
    original_start = workflow.start_segment
    monkeypatch.setattr(
        workflow,
        "start_segment",
        lambda *args, **kwargs: (_ for _ in ()).throw(
            ClientError(
                {"Error": {"Code": "AccessDeniedException", "Message": "synthetic rejection"}},
                "StartExecution",
            )
        ),
    )
    response = client.post(
        f"/v1/jobs/{job_id}/reviews",
        json={
            "expected_revision": store.item["revision"],
            "decisions": {identity: {"decision": "accept"}},
        },
    )
    assert response.status_code == 503
    assert store.item["status"] == "FAILED"
    assert store.item["workflow_segment"] == 2
    assert store.item["error_code"] == "REVIEW_CONTINUATION_FAILED"
    monkeypatch.setattr(workflow, "start_segment", original_start)
    assert client.post(f"/v1/jobs/{job_id}/retry").status_code == 200
    assert workflow.started == [2]


def test_start_metadata_does_not_clobber_worker_status_race(api, monkeypatch):
    client, store, _storage, workflow = api
    job_id = prepare(api)

    def start_and_race(*args, **kwargs):
        store.item["status"] = "DETECTED"
        store.item["revision"] += 1
        return {"execution_arn": "arn:synthetic:execution:race"}

    monkeypatch.setattr(workflow, "start_segment", start_and_race)
    response = client.post(
        f"/v1/jobs/{job_id}/start", params={"expected_revision": store.item["revision"]}
    )
    assert response.status_code == 200, response.text
    assert store.item["status"] == "DETECTED"
    assert store.item["execution_arn"] == "arn:synthetic:execution:race"


def test_review_metadata_does_not_clobber_worker_status_race(api, monkeypatch):
    client, store, _storage, workflow = api
    job_id, identity = ready_quality(api)

    def start_and_race(*args, **kwargs):
        store.item["status"] = "FORMULA_REVIEW_PENDING"
        store.item["revision"] += 1
        return {"execution_arn": "arn:synthetic:execution:review-race"}

    monkeypatch.setattr(workflow, "start_segment", start_and_race)
    response = client.post(
        f"/v1/jobs/{job_id}/reviews",
        json={
            "expected_revision": store.item["revision"],
            "decisions": {identity: {"decision": "accept"}},
        },
    )
    assert response.status_code == 200, response.text
    assert store.item["status"] == "FORMULA_REVIEW_PENDING"
    assert store.item["execution_arn"] == "arn:synthetic:execution:review-race"


def test_retry_started_before_metadata_write_is_not_reported_as_failed(api, monkeypatch):
    client, store, _storage, workflow = api
    job_id, _identity = ready_quality(api)
    store.item.update(status="FAILED", workflow_segment=2, error_code="OLD_FAILURE")

    def start_and_race(*args, **kwargs):
        store.item["status"] = "FORMULA_REVIEW_PENDING"
        store.item["revision"] += 1
        return {"execution_arn": "arn:synthetic:execution:retry-race"}

    monkeypatch.setattr(workflow, "start_segment", start_and_race)
    response = client.post(f"/v1/jobs/{job_id}/retry")
    assert response.status_code == 200, response.text
    assert store.item["status"] == "FORMULA_REVIEW_PENDING"
    assert store.item["execution_arn"] == "arn:synthetic:execution:retry-race"


def test_retry_unknown_start_is_not_retried_as_a_new_execution(api, monkeypatch):
    client, store, _storage, workflow = api
    job_id, _identity = ready_quality(api)
    store.item.update(status="FAILED", workflow_segment=2, error_code="OLD_FAILURE")
    calls = 0

    def unknown_start(*args, **kwargs):
        nonlocal calls
        calls += 1
        raise TimeoutError("response lost after StartExecution")

    monkeypatch.setattr(workflow, "start_segment", unknown_start)
    response = client.post(f"/v1/jobs/{job_id}/retry")
    assert response.status_code == 503
    assert store.item["status"] == "NORMALIZATION_PENDING"
    assert store.item["error_code"] == "WORKFLOW_START_UNKNOWN"
    assert client.post(f"/v1/jobs/{job_id}/retry").status_code == 409
    assert calls == 1


def test_retry_definitive_start_failure_marks_job_failed(api, monkeypatch):
    client, store, _storage, workflow = api
    job_id, _identity = ready_quality(api)
    store.item.update(status="FAILED", workflow_segment=2, error_code="OLD_FAILURE")
    rejection = ClientError(
        {"Error": {"Code": "AccessDeniedException", "Message": "denied"}},
        "StartExecution",
    )
    monkeypatch.setattr(
        workflow,
        "start_segment",
        lambda *args, **kwargs: (_ for _ in ()).throw(rejection),
    )

    response = client.post(f"/v1/jobs/{job_id}/retry")

    assert response.status_code == 503
    assert store.item["status"] == "FAILED"
    assert store.item["error_code"] == "WORKFLOW_START_FAILED"


def test_formula_approval_is_blocked_when_output_column_positions_differ(api):
    client, store, storage, _ = api
    job_id = prepare(api, update=True)
    main_key = store.item["files"]["main"]["key"]
    book = load_workbook(BytesIO(storage.objects[main_key]))
    book["Data"]["B3"] = "Status"
    book["Data"]["C3"] = "Amount"
    stream = BytesIO()
    book.save(stream)
    book.close()
    storage.objects[main_key] = stream.getvalue()
    identity = proposal(
        store,
        storage,
        "tool3",
        {
            "formula_template": "=B{row}*2",
            "target_header": "Double",
            "referenced_headers": [],
            "explanation": "Double amount",
            "confidence": 1,
        },
    )
    handlers.quality_gate_handler({"job_id": job_id, "gate": "transformation"}, None)
    row = next(row for row in report(client, store, "transformation") if row["id"] == identity)
    assert row["formula"]["syntax"] == "invalid"
    assert "accept" not in row["allowedDecisions"]
    assert (
        client.post(
            f"/v1/jobs/{job_id}/reviews",
            json={
                "expected_revision": store.item["revision"],
                "decisions": {identity: {"decision": "accept"}},
            },
        ).status_code
        == 422
    )


def test_blank_row_evidence_handles_empty_cells_without_fabricated_coordinates(api):
    client, store, storage, _workflow = api
    job_id = prepare(api)
    key = store.item["files"]["input"]["key"]
    book = load_workbook(BytesIO(storage.objects[key]))
    book["Data"].insert_rows(5)
    stream = BytesIO()
    book.save(stream)
    book.close()
    storage.objects[key] = stream.getvalue()
    handlers.quality_handler({"job_id": job_id}, None)
    handlers.quality_gate_handler({"job_id": job_id}, None)
    rows = report(client, store, "quality")
    finding = next(row for row in rows if row["kind"] == "blank-row")
    assert finding["row"] == 5
    assert finding["cells"][0]["coordinate"] == "A5"
    assert finding["cells"][0].get("before") is None


def test_formula_templates_use_final_row_positions_after_accepted_deletions():
    from app.workers.apply_changes import apply_workbook_changes

    book = load_workbook(BytesIO(workbook()))
    book["Data"].insert_rows(5)
    stream = BytesIO()
    book.save(stream)
    book.close()
    result, _audit = apply_workbook_changes(
        stream.getvalue(),
        sheet_name="Data",
        header_row=3,
        row_deletions={5},
        formula={
            "formula_template": "=B{row}*2",
            "target_column": 4,
            "target_header": "Double",
            "start_row": 4,
            "end_row": 6,
        },
    )
    output = load_workbook(BytesIO(result), data_only=False)
    assert output["Data"].max_row == 5
    assert output["Data"]["D5"].value == "=B5*2"
    output.close()


def test_unknown_start_outcome_does_not_enable_a_duplicate_execution(api, monkeypatch):
    client, store, _storage, workflow = api
    job_id, identity = ready_quality(api)
    monkeypatch.setattr(
        workflow,
        "start_segment",
        lambda *args, **kwargs: (_ for _ in ()).throw(TimeoutError("synthetic lost response")),
    )
    response = client.post(
        f"/v1/jobs/{job_id}/reviews",
        json={
            "expected_revision": store.item["revision"],
            "decisions": {identity: {"decision": "accept"}},
        },
    )
    assert response.status_code == 503
    assert store.item["status"] == "NORMALIZATION_PENDING"
    assert store.item["error_code"] == "REVIEW_CONTINUATION_UNKNOWN"
    assert client.post(f"/v1/jobs/{job_id}/retry").status_code == 409
    assert report(client, store, "quality")[0]["status"] == "human_accepted"
