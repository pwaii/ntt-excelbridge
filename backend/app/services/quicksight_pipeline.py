"""Flatten per-job S3 artifacts into stable CSV sources for Amazon Quick."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from collections.abc import Iterable, Mapping
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.services.quicksight_identity import reader_user_arn
from app.settings import settings

QUALITY_FIELDS = (
    "job_id",
    "owner_id",
    "input_slot",
    "issue_id",
    "kind",
    "sheet",
    "column",
    "recommended_action",
    "row_count",
    "message",
)
INSPECTION_FIELDS = (
    "job_id",
    "owner_id",
    "input_slot",
    "sheet",
    "total_rows",
    "column_name",
    "column_index",
    "dominant_type",
    "blank_count",
    "blank_rate",
)
PERMISSION_FIELDS = ("UserARN", "owner_id")
WORKBOOK_ROW_FIELDS = (
    "job_id",
    "owner_id",
    "source_kind",
    "input_slot",
    "workbook_name",
    "sheet",
    "row_number",
    "row_data",
    "product_name",
    "product_row",
    "unit_price",
)
_QUALITY_NAME = re.compile(r"(?:(input_[2-8])-)?quality-report\.json\Z")
_INSPECTION_NAME = re.compile(r"(?:(input_[2-8])-)?manifest\.json\Z")
_WORKBOOK_NAME = re.compile(r"(?:input|input_[2-8]|main)/(?:[^/]+)\.xlsx\Z", re.IGNORECASE)
_PRODUCT_HEADER = re.compile(r"\b(product|item|sku|description)\b", re.IGNORECASE)
_PRICE_HEADER = re.compile(r"\b(price|cost|amount|revenue|sales)\b", re.IGNORECASE)
_UNIT_PRICE_HEADER = re.compile(r"\b(unit\s+price|selling\s+price|sale\s+price|list\s+price|price\s+per\s+unit)\b", re.IGNORECASE)


def _job_and_slot(key: str, area: str, name_pattern: re.Pattern[str]) -> tuple[str, str] | None:
    parts = key.split("/")
    if len(parts) != 4 or parts[0] != "jobs" or parts[2] != area:
        return None
    match = name_pattern.fullmatch(parts[3])
    if not match:
        return None
    return parts[1], match.group(1) or "input"


def _report_with_owner(report: tuple[Any, ...]) -> tuple[str, str, Mapping[str, Any]]:
    """Accept the legacy ``(key, report)`` pair and the owner-aware triple."""
    if len(report) == 2:
        key, value = report
        return str(key), "", value
    if len(report) == 3:
        key, owner_id, value = report
        return str(key), str(owner_id or ""), value
    raise ValueError("report entries must contain a key, optional owner, and report")


def flatten_quality_reports(reports: Iterable[tuple[Any, ...]]) -> list[dict[str, Any]]:
    """Return one row per issue, including the source workbook slot."""
    rows: list[dict[str, Any]] = []
    for report_entry in reports:
        key, owner_id, report = _report_with_owner(report_entry)
        parsed = _job_and_slot(key, "quality", _QUALITY_NAME)
        if not parsed:
            continue
        job_id, input_slot = parsed
        for issue in report.get("issues", []):
            rows.append({
                "job_id": job_id,
                "owner_id": owner_id,
                "input_slot": input_slot,
                "issue_id": issue.get("issue_id", ""),
                "kind": issue.get("kind", ""),
                "sheet": issue.get("sheet", ""),
                "column": issue.get("column", ""),
                "recommended_action": issue.get("recommended_action", ""),
                "row_count": len(issue.get("row_numbers", [])),
                "message": issue.get("message", ""),
            })
    return rows


def flatten_inspection_manifests(manifests: Iterable[tuple[Any, ...]]) -> list[dict[str, Any]]:
    """Return one row per workbook column, retaining workbook slot identity."""
    rows: list[dict[str, Any]] = []
    for manifest_entry in manifests:
        key, owner_id, manifest = _report_with_owner(manifest_entry)
        parsed = _job_and_slot(key, "inspection", _INSPECTION_NAME)
        if not parsed:
            continue
        job_id, input_slot = parsed
        total_rows = manifest.get("total_rows", 0)
        for sheet in manifest.get("sheets", []):
            for column in sheet.get("columns", []):
                distribution = column.get("type_distribution", {})
                dominant_type = max(distribution, key=distribution.get) if distribution else "unknown"
                blank_count = column.get("blank_count", 0)
                rows.append({
                    "job_id": job_id,
                    "owner_id": owner_id,
                    "input_slot": input_slot,
                    "sheet": sheet.get("name", ""),
                    "total_rows": total_rows,
                    "column_name": column.get("name", ""),
                    "column_index": column.get("index", ""),
                    "dominant_type": dominant_type,
                    "blank_count": blank_count,
                    "blank_rate": round(blank_count / total_rows, 3) if total_rows else 0,
                })
    return rows


def flatten_workbook_rows(
    workbooks: Iterable[tuple[str, str, bytes]],
    inspection_manifests: Mapping[tuple[str, str], Mapping[str, Any]] | None = None,
) -> list[dict[str, Any]]:
    """Return one owner-filtered record per non-empty spreadsheet data row.

    ``row_data`` stores a readable label/value description using the inspected
    column headers, with detected product and price fields placed first.
    """
    manifests = inspection_manifests or {}
    rows: list[dict[str, Any]] = []
    for key, owner_id, source in workbooks:
        parts = key.split("/")
        if parts[0] != "jobs" or not owner_id:
            continue
        # Input lives under input/<slot>/<file>.xlsx (5 parts); the generated
        # deliverable under output/<file>.xlsx (4 parts, no slot, no manifest).
        if len(parts) == 5 and parts[2] == "input":
            source_kind, job_id, slot, filename = "input", parts[1], parts[3], parts[4]
            if not _WORKBOOK_NAME.fullmatch(f"{slot}/{filename}"):
                continue
        elif len(parts) == 4 and parts[2] == "output" and parts[3].lower().endswith(".xlsx"):
            source_kind, job_id, slot, filename = "output", parts[1], "output", parts[3]
        else:
            continue

        inspection = manifests.get((job_id, slot), {})
        sheet_manifests = {sheet.get("name"): sheet for sheet in inspection.get("sheets", [])}
        workbook = load_workbook(BytesIO(source), read_only=True, data_only=False)
        try:
            for worksheet in workbook.worksheets:
                sheet_manifest = sheet_manifests.get(worksheet.title, {})
                header_row = int(sheet_manifest.get("header_row") or 1)
                headers = list(sheet_manifest.get("headers") or [])
                if not headers:
                    # No manifest (e.g. generated output): read the real header row
                    # instead of fabricating Column_N, so field names stay meaningful.
                    headers = list(next(worksheet.iter_rows(
                        min_row=header_row, max_row=header_row, values_only=True), ()))
                width = max(len(headers), worksheet.max_column or 0)
                headers = [
                    str(headers[index]) if index < len(headers) and headers[index] not in (None, "")
                    else f"Column_{index + 1}"
                    for index in range(width)
                ]
                # Make repeated header labels distinct so no cell values are lost.
                unique_headers: list[str] = []
                seen: dict[str, int] = {}
                for index in range(width):
                    base = headers[index] if index < len(headers) else f"Column_{index + 1}"
                    seen[base] = seen.get(base, 0) + 1
                    unique_headers.append(base if seen[base] == 1 else f"{base}_{seen[base]}")

                for row_number, values in enumerate(
                    worksheet.iter_rows(min_row=header_row + 1, max_col=width, values_only=True),
                    start=header_row + 1,
                ):
                    if row_number > settings.max_data_rows:
                        raise ValueError(f"workbook exceeds the {settings.max_data_rows} row limit")
                    if not any(value is not None and str(value).strip() for value in values):
                        continue
                    row_values: dict[str, Any] = {}
                    row_value_types: dict[str, str] = {}
                    cells = []
                    for index, value in enumerate(values):
                        if value is None or not str(value).strip():
                            row_values[unique_headers[index]] = None
                            continue
                        header = unique_headers[index]
                        cells.append((header, _display_cell(value)))
                        row_values[header] = _json_cell(value)
                        row_value_types[header] = _cell_kind(value)
                    row_data = _describe_cells(cells)
                    product = _product_cell(cells)
                    unit_price = _unit_price_cell(cells)
                    product_name = product[1] if product else ""
                    row_ref = hashlib.sha256(
                        f"{job_id}\0{slot}\0{worksheet.title}\0{row_number}".encode("utf-8")
                    ).hexdigest()[:8]
                    product_row = " · ".join(
                        part for part in (
                            product_name or "Product not identified",
                            filename,
                            worksheet.title,
                            f"row {row_number}",
                            row_ref,
                        ) if part
                    )
                    rows.append({
                        "job_id": job_id,
                        "owner_id": owner_id,
                        "source_kind": source_kind,
                        "input_slot": slot,
                        "workbook_name": filename,
                        "sheet": worksheet.title,
                        "row_number": row_number,
                        "row_data": row_data,
                        "product_name": product_name,
                        "product_row": product_row,
                        "unit_price": unit_price,
                        # Used by the website's schema-driven explorer. QuickSight's
                        # stable CSV schema ignores these additional in-memory keys.
                        "row_values": row_values,
                        "row_value_types": row_value_types,
                    })
        finally:
            workbook.close()
    return rows


def _display_cell(value: Any) -> str:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _json_cell(value: Any) -> Any:
    """Preserve Excel values as JSON-compatible scalars for field selection."""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, (bool, int, float)) or value is None:
        return value
    return str(value).strip()


def _cell_kind(value: Any) -> str:
    if isinstance(value, (date, datetime)):
        return "date"
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        return "number"
    if isinstance(value, str) and _decimal_value(value):
        return "number"
    return "category"


def _describe_cells(cells: list[tuple[str, str]]) -> str:
    """Format a worksheet row as readable header/value pairs."""
    if not cells:
        return ""
    product = _product_cell(cells)
    price = next((cell for cell in cells if _PRICE_HEADER.search(_header_semantics(cell[0]))), None)
    primary = []
    consumed: set[int] = set()
    if product:
        primary.append(f"Product: {product[1]}")
        consumed.add(cells.index(product))
    if price and price != product:
        primary.append(f"{_friendly_label(price[0])}: {price[1]}")
        consumed.add(cells.index(price))
    details = [
        f"{_friendly_label(header)}: {value}"
        for index, (header, value) in enumerate(cells)
        if index not in consumed
    ]
    return " | ".join(primary + details)


def _product_cell(cells: list[tuple[str, str]]) -> tuple[str, str] | None:
    """Pick the most descriptive product/item field from a worksheet row."""
    return next(
        (cell for cell in cells if _PRODUCT_HEADER.search(_header_semantics(cell[0])) and re.search(r"\b(name|title|description)\b", _header_semantics(cell[0]), re.IGNORECASE)),
        None,
    ) or next((cell for cell in cells if _PRODUCT_HEADER.search(_header_semantics(cell[0]))), None)


def _unit_price_cell(cells: list[tuple[str, str]]) -> str:
    """Return the row's unit price as a decimal string, or blank if unavailable.

    Prefer explicit per-unit and selling-price headers. Generic price is the
    fallback. Cost, revenue, sales, and total amount are intentionally excluded
    because they do not reliably mean unit price.
    """
    explicit = next((cell for cell in cells if _UNIT_PRICE_HEADER.search(_header_semantics(cell[0]))), None)
    fallback = next(
        (cell for cell in cells if re.search(r"\bprice\b", _header_semantics(cell[0]), re.IGNORECASE)),
        None,
    )
    candidate = explicit or fallback
    return _decimal_value(candidate[1]) if candidate else ""


def _decimal_value(value: str) -> str:
    """Normalize common spreadsheet price formats without guessing currencies."""
    normalized = value.strip().replace("−", "-")
    is_negative = normalized.startswith("(") and normalized.endswith(")")
    if is_negative:
        normalized = normalized[1:-1]
    normalized = re.sub(r"[,$€£¥₹\s]", "", normalized)
    if normalized.startswith("-"):
        is_negative = not is_negative
        normalized = normalized[1:]
    if not normalized:
        return ""
    if "," in normalized:
        if "." in normalized:
            normalized = normalized.replace(",", "")
        else:
            integer, separator, fractional = normalized.rpartition(",")
            if separator and 1 <= len(fractional) <= 2 and integer:
                normalized = f"{integer.replace(',', '')}.{fractional}"
            else:
                normalized = normalized.replace(",", "")
    if not re.fullmatch(r"\d+(?:\.\d+)?|\.\d+", normalized):
        return ""
    try:
        number = Decimal(normalized)
    except InvalidOperation:
        return ""
    if not number.is_finite():
        return ""
    normalized = format(number, "f")
    if "." in normalized:
        normalized = normalized.rstrip("0").rstrip(".")
    return f"-{normalized}" if is_negative and number else normalized


def _friendly_label(header: str) -> str:
    label = re.sub(r"[_-]+", " ", header).strip()
    if label.casefold() == "sku":
        return "SKU"
    return label[:1].upper() + label[1:]


def _header_semantics(header: str) -> str:
    return re.sub(r"[_-]+", " ", header).strip()


def encode_csv(rows: list[dict[str, Any]], fields: tuple[str, ...]) -> bytes:
    """Create a header-only CSV for empty datasets instead of indexing rows[0]."""
    buffer = io.StringIO(newline="")
    writer = csv.DictWriter(buffer, fieldnames=fields, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8")


def owner_permissions(owner_ids: Iterable[str]) -> list[dict[str, str]]:
    """Build unique QuickSight user-ARN to app-owner RLS grants."""
    if not settings.quicksight_embedding_role_arn:
        return []
    rows = {
        (reader_user_arn(settings.quicksight_embedding_role_arn, settings.quicksight_region, owner_id), owner_id)
        for owner_id in owner_ids
        if owner_id
    }
    return [{"UserARN": user_arn, "owner_id": owner_id} for user_arn, owner_id in sorted(rows)]


def parse_artifact_json(key: str, body: bytes, area: str, name_pattern: re.Pattern[str]) -> tuple[str, dict[str, Any]] | None:
    """Parse an expected JSON artifact and return its identifying metadata."""
    parsed = _job_and_slot(key, area, name_pattern)
    if parsed is None:
        return None
    value = json.loads(body)
    if not isinstance(value, dict):
        raise ValueError(f"Expected an object in s3://.../{key}")
    return parsed[0], value


def quality_artifact_key(key: str) -> bool:
    return _job_and_slot(key, "quality", _QUALITY_NAME) is not None


def inspection_artifact_key(key: str) -> bool:
    return _job_and_slot(key, "inspection", _INSPECTION_NAME) is not None
