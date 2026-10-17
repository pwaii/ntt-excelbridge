"""Read metadata from uploaded bytes; never evaluate formulas or return row data."""

from hashlib import sha256
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.domain.workbooks import WorkbookColumn, WorksheetMetadata
from app.settings import settings


def inspect_upload(source: bytes, sheet_name: str | None = None, header_row: int | None = None):
    if len(source) > settings.max_workbook_mb * 1024 * 1024:
        raise ValueError("Workbook exceeds the file size ceiling")
    try:
        with ZipFile(BytesIO(source)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 256 * 1024 * 1024:
                raise ValueError("Workbook expanded content exceeds the 256 MiB safety limit")
        workbook = load_workbook(BytesIO(source), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, KeyError, OSError) as error:
        raise ValueError("The uploaded object is not a readable XLSX workbook") from error
    try:
        if sheet_name is not None and sheet_name not in workbook.sheetnames:
            raise ValueError("Selected worksheet does not exist")
        sheets = []
        total_rows = 0
        for sheet in workbook.worksheets:
            # Count actual occupied rows rather than trusting XLSX dimensions.
            last = 0
            candidates = []
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                nonempty = [value for value in row if value is not None and str(value).strip()]
                if nonempty:
                    last = row_number
                if row_number <= 20 and nonempty:
                    score = (
                        len(nonempty),
                        sum(isinstance(value, str) for value in nonempty),
                        len(set(map(str, nonempty))),
                        -row_number,
                    )
                    candidates.append((score, row_number, row))
                if row_number > settings.max_data_rows + 20:
                    raise ValueError("Workbook exceeds the 50,000-row ceiling")
            guess = max(candidates, default=None, key=lambda item: item[0])
            detected = guess[1] if guess else None
            chosen = (
                header_row if sheet.title == sheet_name and header_row is not None else detected
            )
            values = (
                next(sheet.iter_rows(min_row=chosen, max_row=chosen, values_only=True), ())
                if chosen
                else ()
            )
            columns = [
                WorkbookColumn(id=get_column_letter(index), label=str(value).strip())
                for index, value in enumerate(values, 1)
                if value is not None and str(value).strip()
            ]
            errors = (
                [] if columns else ["No nonempty headers were found. Select a valid header row."]
            )
            count = max(0, last - (chosen or 0))
            total_rows += count
            if total_rows > settings.max_data_rows:
                raise ValueError("Workbook exceeds the 50,000-row ceiling")
            sheets.append(
                WorksheetMetadata(
                    name=sheet.title,
                    detected_header_row=detected,
                    header_row=chosen,
                    data_row_count=count,
                    columns=columns,
                    errors=errors,
                )
            )
        return sha256(source).hexdigest(), sheets
    finally:
        workbook.close()
