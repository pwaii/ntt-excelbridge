"""Build review evidence from actual owned S3 artifacts, never browser guesses."""

import json
from io import BytesIO

from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string, get_column_letter

from app.domain.review_report import CellEvidence, FormulaEvidence, MatchEvidence, ReportItem
from app.workers.apply_changes import materialize_formula
from app.workers.formulas import verify_formula
from app.workers.inspect_workbook import normalize_text


def all_children(dynamo, job_id, prefix):
    """Consume every DynamoDB page; never silently drop decisions after row 100."""
    records, cursor, seen = [], None, set()
    while True:
        page, following = dynamo.list_children_page(job_id, prefix, limit=100, cursor=cursor)
        records.extend(page)
        if not following:
            return records
        if following in seen:
            raise HTTPException(503, "Report pagination did not advance")
        seen.add(following)
        cursor = following


def _slot_qualified(slot, bare_id):
    """Primary input keeps bare ids (back-compat); extra inputs are slot-prefixed so
    two workbooks that share a sheet name never collide on an issue/proposal id."""
    return bare_id if slot == "input" else f"{slot}::{bare_id}"


def _slot_sheet_header(item, slot):
    """This slot's own (sheet_name, header_row). Primary 'input'/'main' live in config;
    extra inputs keep their selection in workbook_selections. Mirrors lambda_handlers so
    the report never indexes a secondary workbook with the primary input's sheet name."""
    config = item.get("config", {})
    if slot == "input":
        return config.get("input_sheet"), config.get("input_header_row") or 1
    if slot == "main":
        return config.get("main_sheet"), config.get("main_header_row") or 1
    sel = item.get("workbook_selections", {}).get(slot, {})
    return sel.get("sheet_name"), sel.get("header_row") or 1


_QUALITY_KINDS = {
    "duplicate_rows": "duplicate-row",
    "blank_row": "blank-row",
    "blank_column": "blank-column",
    "blank_key_field": "blank-key-field",
    "type_mismatch": "type-mismatch",
}


def _quality_rows(issues, s3, item, slot, filename, prefix):
    """Build the review rows for one input slot's quality issues, with cell evidence
    drawn from that slot's own workbook."""
    workbook = None
    source = item.get("files", {}).get(slot, {}).get("key")
    rows = []
    try:
        if source and source.startswith(prefix):
            workbook = load_workbook(BytesIO(s3.get_bytes(source)), read_only=True, data_only=False)
        requested_rows = {}
        for issue in issues:
            requested_rows.setdefault(issue.get("sheet"), set()).update(issue.get("row_numbers", [])[:3])
        samples = {}
        if workbook:
            for sheet_name, numbers in requested_rows.items():
                if not numbers or sheet_name not in workbook.sheetnames:
                    continue
                for number, values in enumerate(
                    workbook[sheet_name].iter_rows(
                        min_row=min(numbers), max_row=max(numbers), max_col=8, values_only=True
                    ),
                    min(numbers),
                ):
                    if number in numbers:
                        samples[(sheet_name, number)] = values
        for issue in issues:
            if issue.get("kind") not in _QUALITY_KINDS:
                raise HTTPException(503, "Unsupported quality finding")
            kind = _QUALITY_KINDS[issue["kind"]]
            cells = []
            sheet_name = issue.get("sheet")
            row_numbers = issue.get("row_numbers", [])
            if workbook and sheet_name in workbook.sheetnames:
                for number in row_numbers[:3]:
                    for column, value in enumerate(samples.get((sheet_name, number), ()), 1):
                        cells.append(
                            CellEvidence(
                                coordinate=f"{get_column_letter(column)}{number}",
                                before=None if value is None else str(value),
                                after="Keep first row; remove later duplicates"
                                if kind == "duplicate-row"
                                else "No automatic change; fill in the key manually"
                                if kind == "blank-key-field"
                                else "Remove row if accepted",
                            )
                        )
            rows.append(
                ReportItem(
                    id=_slot_qualified(slot, issue["issue_id"]),
                    kind=kind,
                    title=issue["message"],
                    reason=issue["message"]
                    + (f" Rows: {row_numbers[:50]}" if row_numbers else "")
                    + (
                        " Informational finding; no supported direct mutation is proposed."
                        if kind in {"type-mismatch", "blank-column", "blank-key-field"}
                        else ""
                    ),
                    slot=slot,
                    filename=filename,
                    sheet=sheet_name,
                    row=row_numbers[0] if row_numbers else None,
                    column=str(issue["column"]) if issue.get("column") else None,
                    cells=cells,
                    requires_decision=kind not in {"type-mismatch", "blank-column", "blank-key-field"},
                    affected_count=len(row_numbers) or None,
                )
            )
        return rows
    finally:
        if workbook:
            workbook.close()


def build_report(item, stage, dynamo, s3):
    cache = {}
    prefix = s3.job_prefix(item["job_id"]) + "/"

    def artifact(key):
        if not isinstance(key, str) or not key.startswith(prefix):
            raise HTTPException(503, "Review artifact is unavailable")
        if key not in cache:
            data = s3.get_bytes(key)
            if len(data) > 20 * 1024 * 1024:
                raise HTTPException(503, "Review artifact exceeds the supported report size")
            cache[key] = json.loads(data)
        return cache[key]

    rows = []
    config = item.get("config", {})
    source = item.get("files", {}).get("input", {}).get("key")
    workbook = None
    try:
        if source:
            if not source.startswith(prefix):
                raise HTTPException(503, "Review source is unavailable")
            workbook = load_workbook(BytesIO(s3.get_bytes(source)), read_only=True, data_only=False)
        if stage == "quality":
            input_slots = [s for s in (["input"] + [f"input_{i}" for i in range(2, 9)]) if s in item.get("files", {})] or ["input"]
            found_any = False
            for slot in input_slots:
                artifacts_map = item.get("artifacts", {})
                key = artifacts_map.get("quality" if slot == "input" else f"quality_{slot}", {}).get("key")
                if not key:
                    continue
                found_any = True
                filename = item.get("files", {}).get(slot, {}).get("filename")
                rows.extend(_quality_rows(artifact(key).get("issues", []), s3, item, slot, filename, prefix))
            if not found_any:
                raise HTTPException(409, "The quality report is not ready")

        for proposal in all_children(dynamo, item["job_id"], "PROPOSAL#"):
            tool = proposal.get("tool")
            if tool not in (
                {"tool1", "tool2"} if stage == "quality" else {"tool3", "tool4", "tool5"}
            ):
                continue
            payload = artifact(proposal.get("artifact_key"))["payload"]
            proposal_slot = proposal.get("slot", "input")
            proposal_filename = item.get("files", {}).get(proposal_slot, {}).get("filename")
            identity = _slot_qualified(proposal_slot, proposal["proposal_id"])
            index = proposal.get("proposal_index")
            if tool == "tool1":
                rows.append(
                    ReportItem(
                        id=identity,
                        slot=proposal_slot,
                        filename=proposal_filename,
                        kind="semantic",
                        title=payload["meaning"],
                        reason=payload["interpretation"],
                        confidence=payload["confidence"],
                    )
                )
            elif tool == "tool2":
                if not payload.get("rules"):
                    continue
                if index is None or not 0 <= int(index) < len(payload.get("rules", [])):
                    raise HTTPException(503, "Normalization rule identity is unavailable")
                rule = payload["rules"][int(index)]
                if rule["example_original"] == rule["normalized_form"]:
                    continue
                rows.append(
                    ReportItem(
                        id=identity,
                        slot=proposal_slot,
                        filename=proposal_filename,
                        kind="normalization-rule",
                        title=f"Normalize {payload['column_id']}",
                        column=payload["column_id"],
                        reason=rule["reason"],
                        pattern=rule["pattern"],
                        confidence=rule["confidence"],
                        cells=[
                            CellEvidence(
                                coordinate=payload["column_id"],
                                before=rule["example_original"],
                                after=rule["normalized_form"],
                            )
                        ],
                    )
                )
            elif tool == "tool3":
                formula = payload["formula_template"]
                syntax, diagnostic = "unverified", None
                # Use the workbook for this proposal's actual slot (not always the primary input)
                slot_key = item.get("files", {}).get(proposal_slot, {}).get("key")
                if proposal_slot == "input":
                    slot_wb, slot_wb_owned = workbook, False
                elif slot_key and slot_key.startswith(prefix):
                    slot_wb = load_workbook(BytesIO(s3.get_bytes(slot_key)), read_only=True, data_only=False)
                    slot_wb_owned = True
                else:
                    slot_wb, slot_wb_owned = None, False
                slot_sheet, slot_header = _slot_sheet_header(item, proposal_slot)
                try:
                    if slot_wb:
                        sheet = (
                            slot_wb[slot_sheet]
                            if slot_sheet
                            else slot_wb.worksheets[0]
                        )
                        header = slot_header
                        headers = [cell.value for cell in sheet[header]]
                        # Normalize for comparison — inspect_bytes normalizes header names
                        headers_normalized = {normalize_text(h) for h in headers if h is not None}
                        target = payload.get("target_header") or config.get("output_header")
                        target_column = (
                            headers.index(target) + 1 if target in headers else sheet.max_column + 1
                        )
                        try:
                            if any(name not in headers_normalized for name in payload["referenced_headers"]):
                                raise ValueError(
                                    "Formula references a header outside the confirmed source"
                                )
                            if config.get("update"):
                                main_key = item.get("files", {}).get("main", {}).get("key")
                                if not main_key or not main_key.startswith(prefix):
                                    raise ValueError(
                                        "The main workbook is unavailable for formula verification"
                                    )
                                main_book = load_workbook(
                                    BytesIO(s3.get_bytes(main_key)), read_only=True, data_only=False
                                )
                                try:
                                    main_sheet = (
                                        main_book[config["main_sheet"]]
                                        if config.get("main_sheet")
                                        else main_book.worksheets[0]
                                    )
                                    main_headers = [
                                        cell.value
                                        for cell in main_sheet[config.get("main_header_row") or 1]
                                    ]
                                    if any(
                                        name not in main_headers
                                        or main_headers.index(name) != headers.index(name)
                                        for name in payload["referenced_headers"]
                                    ):
                                        raise ValueError(
                                            "Referenced columns differ in the output workbook; request a new output-specific formula"
                                        )
                                    main_header = config.get("main_header_row") or 1
                                    main_target = (
                                        main_headers.index(target) + 1
                                        if target in main_headers
                                        else main_sheet.max_column + 1
                                    )
                                    if main_sheet.max_row <= main_header:
                                        raise ValueError(
                                            "The main sheet has no output rows for this formula"
                                        )
                                    for output_row in {main_header + 1, main_sheet.max_row}:
                                        verified = verify_formula(
                                            materialize_formula(formula, output_row),
                                            min_row=main_header + 1,
                                            max_row=main_sheet.max_row,
                                            max_column=max(main_sheet.max_column, main_target),
                                            target_column=main_target,
                                        )
                                        for reference in verified["references"]:
                                            letter, _number = coordinate_from_string(
                                                reference.replace("$", "")
                                            )
                                            position = column_index_from_string(letter) - 1
                                            if (
                                                position >= len(headers)
                                                or position >= len(main_headers)
                                                or headers[position] != main_headers[position]
                                            ):
                                                raise ValueError(
                                                    "Formula cell references differ in the output workbook"
                                                )
                                finally:
                                    main_book.close()
                            if sheet.max_row <= header:
                                raise ValueError("The selected sheet has no data rows")
                            for row in {header + 1, sheet.max_row}:
                                verify_formula(
                                    materialize_formula(formula, row),
                                    min_row=header + 1,
                                    max_row=sheet.max_row,
                                    max_column=max(sheet.max_column, target_column),
                                    target_column=target_column,
                                )
                            syntax = "verified"
                        except ValueError as error:
                            syntax, diagnostic = "invalid", str(error)
                finally:
                    if slot_wb_owned:
                        slot_wb.close()
                rows.append(
                    ReportItem(
                        id=identity,
                        slot=proposal_slot,
                        filename=proposal_filename,
                        kind="formula",
                        title=f"Formula: {payload['target_header']}",
                        reason=payload["explanation"],
                        confidence=payload["confidence"],
                        formula=FormulaEvidence(
                            text=formula,
                            referenced_columns=payload["referenced_headers"],
                            syntax=syntax,
                            diagnostic=diagnostic,
                        ),
                        allowed_decisions=["accept", "reject", "ignore"]
                        if syntax == "verified"
                        else ["reject", "ignore"],
                    )
                )
            elif tool == "tool4":
                mappings = payload.get("mappings", [])
                # Legacy grouped proposals retain their actual parent decision ID.
                selected = [mappings[int(index)]] if index is not None else mappings
                rows.append(
                    ReportItem(
                        id=identity,
                        slot=proposal_slot,
                        filename=proposal_filename,
                        kind="standardization",
                        title="Standardization mapping"
                        if index is not None
                        else "Standardization mapping group",
                        reason="; ".join(mapping["reason"] for mapping in selected)
                        or "No mappings proposed",
                        confidence=min(
                            (mapping["confidence"] for mapping in selected), default=None
                        ),
                        cells=[
                            CellEvidence(
                                coordinate=mapping["source_field"],
                                before=str(mapping["source_value"]),
                                after=str(mapping["canonical_value"]),
                            )
                            for mapping in selected
                        ],
                        requires_decision=bool(selected),
                        allowed_decisions=["accept", "reject", "edit", "ignore"]
                        if len(selected) == 1
                        else ["accept", "reject", "ignore"],
                    )
                )
            elif tool == "tool5":
                matches = payload.get("matches", [])
                if not matches:
                    continue
                if index is None or not 0 <= int(index) < len(matches):
                    raise HTTPException(503, "Update match identity is unavailable")
                match = matches[int(index)]
                match_type = match["match_type"]
                allowed = (
                    ["reject", "ignore"]
                    if match_type == "unmatched"
                    else ["reject", "edit", "ignore"]
                    if match_type in {"ambiguous", "duplicate", "conflict"}
                    else ["accept", "reject", "edit", "ignore"]
                )
                rows.append(
                    ReportItem(
                        id=identity,
                        slot=proposal_slot,
                        filename=proposal_filename,
                        kind="update-match",
                        title=f"{match_type}: {match['current_row_id']}",
                        reason=match["reason"],
                        confidence=match["confidence"],
                        match=MatchEvidence(
                            type=match_type,
                            source_key=match["current_row_id"],
                            target_keys=match["candidate_main_row_ids"],
                        ),
                        requires_decision=match_type != "unmatched",
                        allowed_decisions=allowed,
                    )
                )
        decisions = {}
        for reference in sorted(
            item.get("review_artifacts", {}).values(), key=lambda entry: int(entry["revision"])
        ):
            decisions.update(artifact(reference["key"]).get("decisions", {}))
        for review in sorted(
            all_children(dynamo, item["job_id"], "REVIEW#"),
            key=lambda entry: int(entry.get("revision", 0)),
        ):
            if review.get("artifact_key"):
                decisions.update(artifact(review["artifact_key"]).get("decisions", {}))
        statuses = {
            "accept": "human_accepted",
            "edit": "human_edited",
            "reject": "human_rejected",
            "ignore": "human_ignored",
        }
        for row in rows:
            if row.id in decisions:
                row.status = statuses[decisions[row.id]["decision"]]
                row.requires_decision = False
        if len({row.id for row in rows}) != len(rows):
            raise HTTPException(503, "Review identities are duplicated")
        return sorted(rows, key=lambda row: row.id)
    finally:
        if workbook:
            workbook.close()


def validate_decisions(item, payload, rows):
    if int(item["revision"]) != payload.expected_revision:
        raise HTTPException(409, "REVIEW_REVISION_CONFLICT: reload before submitting")
    required = {row.id: row for row in rows if row.requires_decision and row.status == "proposed"}
    if set(payload.decisions) != set(required):
        raise HTTPException(422, "Decide exactly all required items in the current gate")
    for identity, draft in payload.decisions.items():
        row = required[identity]
        if draft.decision.value not in row.allowed_decisions:
            raise HTTPException(422, "A decision is not allowed for this proposal")
        if draft.decision.value == "ignore" and not (draft.reason or "").strip():
            raise HTTPException(422, "Ignoring a finding requires a reason")
        if draft.decision.value == "edit":
            if draft.edited_value is None:
                raise HTTPException(422, "An edited value is required")
            if row.match and draft.edited_value not in row.match.target_keys:
                raise HTTPException(422, "Choose one of the supplied target row IDs")
